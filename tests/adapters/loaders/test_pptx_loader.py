"""Tests for the PPTX loader.

Decks are built with ``python-pptx``, so each test states the slide shape it is
exercising rather than hiding it in a committed binary.
"""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from pydantic_rag_agent.adapters.loaders.pptx import PPTX_MEDIA_TYPE, PptxLoader

DOCUMENT_ID = "e" * 64


def build_deck(path: Path, *, title: str | None = None, notes: str | None = None) -> Path:
    """A two-slide deck: one with text and a table, one with a picture."""
    deck = Presentation()
    if title:
        deck.core_properties.title = title

    first = deck.slides.add_slide(deck.slide_layouts[1])
    first.shapes.title.text = "Backups"
    body = first.placeholders[1].text_frame
    body.text = "Run nightly at 02:00 UTC."
    body.add_paragraph().text = "Retained for 30 days."

    table = first.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(3), Inches(1)).table
    table.cell(0, 0).text = "host"
    table.cell(0, 1).text = "role"
    table.cell(1, 0).text = "db-1"
    table.cell(1, 1).text = "primary"

    if notes:
        first.notes_slide.notes_text_frame.text = notes

    second = deck.slides.add_slide(deck.slide_layouts[5])
    second.shapes.title.text = "Topology"
    picture = path.parent / "topology.png"
    Image.new("RGB", (30, 20), "navy").save(picture)
    second.shapes.add_picture(str(picture), Inches(2), Inches(2), Inches(1), Inches(0.7))

    deck.save(path)
    return path


def pptx_loader_should_declare_the_pptx_suffix() -> None:
    # given / when / then: format discovery comes from the loader
    assert PptxLoader().suffixes == frozenset({".pptx"})


def pptx_loader_with_a_deck_should_create_a_block_for_each_slide_with_text(
    tmp_path: Path,
) -> None:
    # given: a two-slide deck, the second holding only a picture
    path = build_deck(tmp_path / "deck.pptx")

    # when: it is loaded
    document = PptxLoader().load(path, document_id=DOCUMENT_ID)

    # then: the text slide becomes a block, and its slide number is preserved
    assert [block.slide for block in document.blocks] == [1]


def pptx_loader_should_use_the_slide_title_as_the_heading_path(tmp_path: Path) -> None:
    # given: slides with titles
    path = build_deck(tmp_path / "deck.pptx")

    # when / then: the title is structure, so a citation can name the slide
    document = PptxLoader().load(path, document_id=DOCUMENT_ID)
    assert [block.heading_path for block in document.blocks] == [("Backups",)]


def pptx_loader_with_a_picture_only_slide_should_carry_its_title_on_the_image(
    tmp_path: Path,
) -> None:
    # given: a slide with a title and nothing but a picture
    path = build_deck(tmp_path / "deck.pptx")

    # when: the deck is loaded
    document = PptxLoader().load(path, document_id=DOCUMENT_ID)

    # then: the image keeps the slide title, so its caption is not orphaned
    assert document.images[0].heading_path == ("Topology",)
    assert document.images[0].slide == 2


def pptx_loader_should_not_repeat_the_title_in_the_body(tmp_path: Path) -> None:
    # given: a slide whose title is also visible text on the slide
    path = build_deck(tmp_path / "deck.pptx")

    # when: the slide body is extracted
    body = PptxLoader().load(path, document_id=DOCUMENT_ID).blocks[0].text

    # then: the title appears as the heading path, not again in the text
    assert "Backups" not in body
    assert "Run nightly at 02:00 UTC." in body


def pptx_loader_should_include_speaker_notes(tmp_path: Path) -> None:
    # given: a slide with notes
    path = build_deck(tmp_path / "deck.pptx", notes="Mention the audit requirement.")

    # when / then: notes are searchable, since they hold what the slide only gestures at
    assert (
        "Mention the audit requirement."
        in PptxLoader().load(path, document_id=DOCUMENT_ID).blocks[0].text
    )


def pptx_loader_with_a_table_should_include_its_cells(tmp_path: Path) -> None:
    # given: a slide containing a table
    path = build_deck(tmp_path / "deck.pptx")

    # when: the slide text is extracted
    body = PptxLoader().load(path, document_id=DOCUMENT_ID).blocks[0].text

    # then: rows survive as pipe-separated lines
    assert "host | role" in body
    assert "db-1 | primary" in body


def pptx_loader_with_a_picture_should_discover_it(tmp_path: Path) -> None:
    # given: a slide with an image
    path = build_deck(tmp_path / "deck.pptx")

    # when: the deck is loaded
    document = PptxLoader().load(path, document_id=DOCUMENT_ID)

    # then: the image is surfaced with content-hash identity and its slide number
    assert len(document.images) == 1
    image = document.images[0]
    assert image.media_type == "image/png"
    assert image.image_id == hashlib.sha256(image.data).hexdigest()
    assert image.slide == 2


def pptx_loader_should_prefer_the_core_properties_title(tmp_path: Path) -> None:
    # given: a deck with an explicit title property
    path = build_deck(tmp_path / "file-name.pptx", title="Operations Handbook")

    # when / then: the property wins over the first slide title
    assert PptxLoader().load(path, document_id=DOCUMENT_ID).title == "Operations Handbook"


def pptx_loader_without_a_property_title_should_use_the_first_slide_title(
    tmp_path: Path,
) -> None:
    # given: a deck with no title property
    path = build_deck(tmp_path / "file-name.pptx")

    # when / then: the first slide's title is the next best thing
    assert PptxLoader().load(path, document_id=DOCUMENT_ID).title == "Backups"


def pptx_loader_should_report_the_pptx_media_type(tmp_path: Path) -> None:
    # given: a valid deck
    path = build_deck(tmp_path / "deck.pptx")

    # when / then: the media type travels with the document
    assert PptxLoader().load(path, document_id=DOCUMENT_ID).media_type == PPTX_MEDIA_TYPE


def pptx_loader_with_a_corrupt_file_should_raise(tmp_path: Path) -> None:
    # given: a file that is not a presentation
    path = tmp_path / "broken.pptx"
    path.write_bytes(b"this is not a pptx")

    # when / then: the loader raises, and the pipeline reports it per file
    with pytest.raises(Exception):  # noqa: B017 - the library picks the exception type
        PptxLoader().load(path, document_id=DOCUMENT_ID)
