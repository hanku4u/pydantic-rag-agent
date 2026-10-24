"""Tests for the DOCX loader.

Fixtures are built with ``python-docx`` itself, which is far better than committing
opaque binaries — each test states exactly the document shape it is testing.
"""

from __future__ import annotations

import hashlib
from pathlib import Path

import docx
import pytest
from docx.opc.exceptions import PackageNotFoundError
from PIL import Image

from pydantic_rag_agent.adapters.loaders.docx import DOCX_MEDIA_TYPE, DocxLoader

DOCUMENT_ID = "d" * 64


def build_handbook(path: Path, *, title: str | None = "Operations Handbook") -> Path:
    """A two-section document with a table in the second section."""
    document = docx.Document()
    if title is not None:
        document.core_properties.title = title

    document.add_heading("Operations", level=1)
    document.add_paragraph("Welcome to the handbook.")
    document.add_heading("Backups", level=2)
    document.add_paragraph("Backups run nightly at 02:00 UTC.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "host"
    table.cell(0, 1).text = "role"
    table.cell(1, 0).text = "db-1"
    table.cell(1, 1).text = "primary"
    document.add_paragraph("It must be rotated quarterly.")

    document.save(path)
    return path


def docx_loader_should_declare_the_docx_suffix() -> None:
    # given / when / then: format discovery comes from the loader
    assert DocxLoader().suffixes == frozenset({".docx"})


def docx_loader_with_headed_sections_should_create_one_block_per_section(
    tmp_path: Path,
) -> None:
    # given: a document with an H1 and an H2
    path = build_handbook(tmp_path / "handbook.docx")

    # when: it is loaded
    document = DocxLoader().load(path, document_id=DOCUMENT_ID)

    # then: heading styles become structure, not body text
    assert [block.heading_path for block in document.blocks] == [
        ("Operations",),
        ("Operations", "Backups"),
    ]
    assert document.blocks[0].text == "Welcome to the handbook."


def docx_loader_with_a_table_should_include_its_cells(tmp_path: Path) -> None:
    # given: a document with a table
    path = build_handbook(tmp_path / "handbook.docx")

    # when: it is loaded
    text = DocxLoader().load(path, document_id=DOCUMENT_ID).blocks[1].text

    # then: rows survive as pipe-separated lines, keeping row association
    assert "host | role" in text
    assert "db-1 | primary" in text


def docx_loader_should_keep_a_table_in_document_order(tmp_path: Path) -> None:
    # given: prose, then a table, then more prose in the same section
    path = build_handbook(tmp_path / "handbook.docx")

    # when: it is loaded
    text = DocxLoader().load(path, document_id=DOCUMENT_ID).blocks[1].text

    # then: the table has not been moved to the end of the section
    assert text.index("nightly") < text.index("db-1") < text.index("rotated")


def docx_loader_should_prefer_the_core_properties_title(tmp_path: Path) -> None:
    # given: a document with an explicit title property
    path = build_handbook(tmp_path / "file-name.docx", title="Operations Handbook")

    # when / then: the property wins over both heading and filename
    assert DocxLoader().load(path, document_id=DOCUMENT_ID).title == "Operations Handbook"


def docx_loader_without_a_property_title_should_use_the_first_heading(
    tmp_path: Path,
) -> None:
    # given: a document with headings but no title property
    path = build_handbook(tmp_path / "file-name.docx", title=None)

    # when / then: the first heading is the next best thing
    assert DocxLoader().load(path, document_id=DOCUMENT_ID).title == "Operations"


def docx_loader_without_headings_or_title_should_use_the_filename(tmp_path: Path) -> None:
    # given: a plain document with neither
    path = tmp_path / "memo.docx"
    document = docx.Document()
    document.add_paragraph("just some prose")
    document.save(path)

    # when / then: the filename is better than no title at all
    loaded = DocxLoader().load(path, document_id=DOCUMENT_ID)
    assert loaded.title == "memo"
    assert loaded.blocks[0].heading_path == ()


def docx_loader_with_an_embedded_image_should_discover_it(tmp_path: Path) -> None:
    # given: a document containing a PNG
    picture = tmp_path / "picture.png"
    Image.new("RGB", (40, 30), "blue").save(picture)

    path = tmp_path / "illustrated.docx"
    document = docx.Document()
    document.add_paragraph("Before the image.")
    document.add_picture(str(picture))
    document.save(path)

    # when: it is loaded
    loaded = DocxLoader().load(path, document_id=DOCUMENT_ID)

    # then: the image is surfaced with content-hash identity, ready for IMG-1
    assert len(loaded.images) == 1
    image = loaded.images[0]
    assert image.media_type == "image/png"
    assert image.image_id == hashlib.sha256(image.data).hexdigest()
    assert image.data


def docx_loader_should_report_the_docx_media_type(tmp_path: Path) -> None:
    # given: a valid document
    path = build_handbook(tmp_path / "handbook.docx")

    # when / then: the media type travels with the document
    assert DocxLoader().load(path, document_id=DOCUMENT_ID).media_type == DOCX_MEDIA_TYPE


def docx_loader_with_a_corrupt_file_should_raise(tmp_path: Path) -> None:
    # given: a file that is not a zip container at all
    path = tmp_path / "broken.docx"
    path.write_bytes(b"this is not a docx")

    # when / then: the loader raises, and the pipeline reports it per file
    with pytest.raises(PackageNotFoundError):
        DocxLoader().load(path, document_id=DOCUMENT_ID)
