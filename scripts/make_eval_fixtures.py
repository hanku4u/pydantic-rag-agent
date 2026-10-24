#!/usr/bin/env python3
"""Generate the binary fixture in the evaluation corpus.

A ``.docx`` is opaque in review, so the file is produced by this script rather than
hand-edited: the script states exactly what the document contains, including the text
that exists only inside its image. Regenerate with:

    uv run python scripts/make_eval_fixtures.py

The point of the image is a fact that appears *nowhere* else in the corpus, so the
evaluation can only answer that question if the caption pipeline worked.
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

import docx
from PIL import Image, ImageDraw

CORPUS = Path("eval/corpus")
IMAGE_TEXT = "ASSET REGISTER: SERVER SRV-042 IS IN RACK B7"


def build_asset_inventory(target: Path) -> None:
    """Write a DOCX whose only mention of rack B7 is inside a picture."""
    with tempfile.TemporaryDirectory() as temp_dir:
        picture = Path(temp_dir) / "rack-label.png"
        image = Image.new("RGB", (640, 140), "white")
        ImageDraw.Draw(image).text((12, 50), IMAGE_TEXT, fill="black")
        image.save(picture)

        document = docx.Document()
        document.core_properties.title = "Asset Inventory"
        document.add_heading("Hardware", level=1)
        document.add_paragraph(
            "Servers are tagged with an asset identifier and a physical location label. "
            "The label is a photograph of the rack plate and is stored alongside the record."
        )
        document.add_picture(str(picture))
        document.add_paragraph(
            "Report a missing or damaged label to the platform team; do not relocate "
            "hardware before the register has been updated."
        )
        document.save(target)


def main() -> int:
    target = CORPUS / "asset-inventory.docx"
    target.parent.mkdir(parents=True, exist_ok=True)
    build_asset_inventory(target)
    print(f"wrote {target} ({target.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
