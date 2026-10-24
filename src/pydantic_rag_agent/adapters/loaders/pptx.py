"""PPTX loading.

One block per slide that has text, with the slide title as the heading path and the
speaker notes appended — notes frequently hold the explanation the slide itself only
gestures at, and they are invisible to anyone reading the deck as a document.

A slide containing only a title and a picture produces no text block; the picture
becomes an image chunk, and it carries the slide title in its heading path so the
caption is not orphaned.

Text frames are walked in shape order. The title placeholder is skipped because it
is already the heading path; including it would prefix every chunk with its own
heading twice.
"""

from __future__ import annotations

import hashlib
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.presentation import Presentation as PresentationType
from pptx.slide import Slide

from pydantic_rag_agent.adapters.loaders._tables import render_table
from pydantic_rag_agent.domain.extraction import ExtractedBlock, ExtractedDocument, ExtractedImage

PPTX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


class PptxLoader:
    """Loads ``.pptx`` files."""

    suffixes = frozenset({".pptx"})

    def load(self, path: Path, *, document_id: str) -> ExtractedDocument:
        presentation = Presentation(str(path))
        blocks, images = _extract(presentation)

        return ExtractedDocument(
            document_id=document_id,
            title=_title(presentation, path),
            source_path=str(path),
            media_type=PPTX_MEDIA_TYPE,
            blocks=tuple(blocks),
            images=tuple(images),
        )


def _title(presentation: PresentationType, path: Path) -> str:
    """Document properties first, then the first slide title, then the filename."""
    properties_title = (presentation.core_properties.title or "").strip()
    if properties_title:
        return properties_title

    for slide in presentation.slides:
        title = _slide_title(slide)
        if title:
            return title

    return path.stem


def _slide_title(slide: Slide) -> str | None:
    title_shape = slide.shapes.title
    if title_shape is None or not title_shape.has_text_frame:
        return None
    text = title_shape.text_frame.text.strip()
    return text or None


def _is_title_placeholder(shape: object) -> bool:
    return bool(
        getattr(shape, "is_placeholder", False)
        and shape.placeholder_format.type  # type: ignore[attr-defined]
        in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
    )


def _extract(
    presentation: PresentationType,
) -> tuple[list[ExtractedBlock], list[ExtractedImage]]:
    blocks: list[ExtractedBlock] = []
    images: list[ExtractedImage] = []
    seen_images: set[str] = set()

    for index, slide in enumerate(presentation.slides, start=1):
        title = _slide_title(slide)
        parts: list[str] = []

        for shape in slide.shapes:
            if _is_title_placeholder(shape):
                continue

            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
            elif getattr(shape, "has_table", False) and shape.has_table:
                rendered = render_table(
                    [cell.text for cell in row.cells] for row in shape.table.rows
                )
                if rendered:
                    parts.append(rendered)
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = _picture(
                    shape,
                    slide_number=index,
                    heading_path=(title,) if title else (),
                )
                if image is not None and image.image_id not in seen_images:
                    seen_images.add(image.image_id)
                    images.append(image)

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Speaker notes: {notes}")

        if parts:
            blocks.append(
                ExtractedBlock(
                    text="\n\n".join(parts),
                    slide=index,
                    heading_path=(title,) if title else (),
                )
            )

    return blocks, images


def _picture(
    shape: object,
    *,
    slide_number: int,
    heading_path: tuple[str, ...],
) -> ExtractedImage | None:
    image = shape.image  # type: ignore[attr-defined]
    data = image.blob
    if not data:
        return None

    return ExtractedImage(
        image_id=hashlib.sha256(data).hexdigest(),
        data=data,
        media_type=image.content_type or "application/octet-stream",
        heading_path=heading_path,
        slide=slide_number,
    )
